"""
Generador de informes ejecutivos PowerPoint basados en la plantilla corporativa Orange.
Feature 010: 010-incident-executive-report
"""

import re
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional, Tuple

import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

try:
    from .executive_models import (
        ExecutiveIncidentData,
        ExecutiveActionPoint,
        ExecutiveTimelineEvent,
        ReportMetadata,
    )
    from .executive_paths import (
        get_executive_template_path,
        get_executive_report_path,
        get_executive_report_filename,
    )
except (ImportError, ValueError):
    from executive_models import (
        ExecutiveIncidentData,
        ExecutiveActionPoint,
        ExecutiveTimelineEvent,
        ReportMetadata,
    )
    from executive_paths import (
        get_executive_template_path,
        get_executive_report_path,
        get_executive_report_filename,
    )


def extract_business_areas(business_impact: str) -> List[str]:
    """
    Extrae los nombres clave de áreas de negocio afectadas (ej. 'Ventas', 'Provisión', 'Atención al Cliente')
    a partir del texto de impacto en negocio, omitiendo descripciones operacionales largas.
    """
    if not business_impact or business_impact.strip() in ["—", "-", "N/A", "None"]:
        return []

    cleaned = re.sub(r'^(?:###?\s*)?Impacto\s+(?:en|de)\s+Negocio\s*[:\n]?', '', business_impact.strip(), flags=re.I)
    cleaned = cleaned.replace('**', '').replace('__', '')

    areas: List[str] = []
    matches = re.findall(
        r'(?:^|[\n\r•\-\*]|(?:Afecci[oó]n\s+(?:a|en)|Afectaci[oó]n\s+(?:a|en)|Impacto\s+en))\s*([A-Za-zÀ-ÿ\s/()]{2,40}?)\s*:\s*',
        cleaned,
        flags=re.I
    )
    for a in matches:
        a = a.strip()
        a = re.sub(r'^(?:Afecci[oó]n\s+(?:a|en)|Afectaci[oó]n\s+(?:a|en)|Impacto\s+en)\s+', '', a, flags=re.I).strip()
        if a and a.lower() not in ['negocio', 'impacto', 'clientes', 'general', 'servicio', 'servicios', 'nota'] and a not in areas:
            areas.append(a)

    if not areas:
        parts = [p.strip().rstrip('.') for p in re.split(r'[,;]\s*|\s+y\s+', cleaned) if p.strip()]
        cleaned_parts = []
        for p in parts:
            p_clean = re.sub(r'^(?:Afecci[oó]n\s+(?:a|en)|Afectaci[oó]n\s+(?:a|en)|Impacto\s+en)\s+', '', p, flags=re.I).strip()
            if p_clean and p_clean.lower() not in ['negocio', 'impacto'] and len(p_clean) < 40:
                cleaned_parts.append(p_clean)
        if cleaned_parts:
            areas = cleaned_parts

    return areas


def format_impact_content(impact_raw: str, business_impact_raw: str = "") -> Tuple[str, str]:
    """
    Formatea la sección de impacto para la presentación ejecutiva:
    - main_impact: Texto completo del impacto general (sin truncar).
    - business_summary: Referencia concisa de áreas afectadas (ej: 'Impacto en Negocio: Ventas, Provisión y Facturación.').
    """
    if not impact_raw:
        return ("Detalle de impacto no disponible.", "")

    # Retrocompatibilidad: si el impacto de negocio venía embebido dentro de impact_raw
    split_m = re.split(r'(?:###?\s*)?Impacto\s+(?:en|de)\s+Negocio\s*[:\n]?', impact_raw, flags=re.I)
    main_text = split_m[0].strip()
    embedded_biz = split_m[1].strip() if len(split_m) > 1 else ""

    main_clean = re.sub(r'^#{1,6}\s+.*$', '', main_text, flags=re.MULTILINE)
    main_lines = [l.strip() for l in main_clean.splitlines() if l.strip()]
    main_impact = " ".join(main_lines) if main_lines else (main_text or "Detalle de impacto no disponible.")

    biz_source = business_impact_raw if (business_impact_raw and business_impact_raw.strip() not in ["—", "-", "N/A"]) else embedded_biz
    areas = extract_business_areas(biz_source)

    business_summary = ""
    if areas:
        if len(areas) > 4:
            areas_str = ", ".join(areas[:3]) + ", etc."
        elif len(areas) > 1:
            areas_str = ", ".join(areas[:-1]) + " y " + areas[-1]
        else:
            areas_str = areas[0]
        business_summary = f"Impacto en Negocio: {areas_str}."

    return (main_impact, business_summary)


def _set_cell_text(
    cell,
    text: str,
    font_size_pt: float = 8.0,
    bold: bool = False,
    color_rgb: RGBColor = RGBColor(0x1A, 0x1A, 0x1A),
):
    """Escribe texto en una celda de tabla con formato explícito garantizando legibilidad."""
    cell.text_frame.clear()
    p = cell.text_frame.paragraphs[0]
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    r = p.add_run()
    r.text = text
    r.font.name = "Calibri"
    r.font.size = Pt(font_size_pt)
    r.font.bold = bold
    r.font.color.rgb = color_rgb


class ExecutiveReportBuilder:
    """Construye y rellena la presentación ejecutiva .pptx a partir de datos estructurados."""

    def __init__(self, template_path: Optional[Path] = None):
        self.template_path = Path(template_path) if template_path else get_executive_template_path()

    @staticmethod
    def _delete_slide(prs, slide_index: int):
        """Elimina limpiamente una diapositiva desvinculando la relación XML del package."""
        slide_id_list = prs.slides._sldIdLst
        if slide_index < len(slide_id_list):
            slide_elem = slide_id_list[slide_index]
            r_id = slide_elem.rId
            try:
                prs.part.drop_rel(r_id)
            except Exception:
                pass
            slide_id_list.remove(slide_elem)

    def generate(self, data: ExecutiveIncidentData, output_path: Optional[Path] = None) -> ReportMetadata:
        """
        Carga la plantilla PowerPoint, rellena las diapositivas con los datos del postmortem
        y guarda el informe en la ruta especificada.
        """
        if not self.template_path.is_file():
            raise FileNotFoundError(f"Plantilla base PowerPoint no encontrada: {self.template_path}")

        prs = pptx.Presentation(str(self.template_path))

        # Título estándar reutilizado en todas las diapositivas
        title_text = f"RESUMEN EJECUTIVO INCIDENCIA- TT {data.incident_ref} - {data.title}"

        # 1. Rellenar Diapositiva 1 (Resumen, impacto, causa, solución y acciones)
        self._populate_slide_1(prs.slides[0], data, title_text)

        # 2. Rellenar Diapositivas de cronología (o eliminarlas si no hay eventos)
        self._populate_timeline_slides(prs, data.timeline_events, title_text)

        # Guardar archivo con tolerancia a bloqueos de Windows (PermissionError)
        dest_path = Path(output_path) if output_path else get_executive_report_path(data.incident_ref)
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            prs.save(str(dest_path))
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            fallback_name = f"{dest_path.stem}_{timestamp}{dest_path.suffix}"
            dest_path = dest_path.parent / fallback_name
            prs.save(str(dest_path))
            print(f"[AVISO] El fichero estaba en uso por otra aplicación. Se guardó como: {dest_path.name}")

        size_bytes = dest_path.stat().st_size
        generated_at = datetime.now(timezone.utc).isoformat()

        return ReportMetadata(
            incident_ref=data.incident_ref,
            filename=dest_path.name,
            file_path=str(dest_path),
            generated_at=generated_at,
            size_bytes=size_bytes,
            slide_count=len(prs.slides),
        )

    def _populate_slide_1(self, slide, data: ExecutiveIncidentData, title_text: str):
        """Cumplimenta las formas de texto y la tabla de acciones en la primera diapositiva."""
        for shape in slide.shapes:
            # 1. Título principal (blanco sobre fondo oscuro corporativo)
            if shape.shape_id == 4 or (shape.has_text_frame and "RESUMEN EJECUTIVO INCIDENCIA" in shape.text_frame.text):
                tf = shape.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.space_after = Pt(0)
                run = p.add_run()
                run.text = title_text
                run.font.name = "Calibri"
                run.font.bold = True
                run.font.size = Pt(10.0) if len(title_text) > 80 else Pt(11.5)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            # 2. Subtítulo (Inicio y Duración)
            elif shape.shape_id == 3 or (shape.has_text_frame and "Inicio:" in shape.text_frame.text):
                tf = shape.text_frame
                tf.clear()
                start_display = data.start_time or "No especificado"
                if "T" in start_display and ("+" in start_display or "Z" in start_display):
                    try:
                        clean_t = start_display.split(".")[0].replace("Z", "")
                        dt = datetime.fromisoformat(clean_t)
                        start_display = dt.strftime("%d/%m/%Y %H:%M")
                    except Exception:
                        pass
                dur_display = data.duration or "No especificada"

                p0 = tf.paragraphs[0]
                p0.space_after = Pt(1)
                r0 = p0.add_run()
                r0.text = "Inicio: "
                r0.font.name = "Calibri"
                r0.font.size = Pt(9.0)
                r0.font.bold = True
                r0.font.color.rgb = RGBColor(0xFF, 0x79, 0x00)
                r1 = p0.add_run()
                r1.text = start_display
                r1.font.name = "Calibri"
                r1.font.size = Pt(9.0)
                r1.font.bold = False
                r1.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

                p1 = tf.add_paragraph()
                p1.space_after = Pt(0)
                r2 = p1.add_run()
                r2.text = "Duración: "
                r2.font.name = "Calibri"
                r2.font.size = Pt(9.0)
                r2.font.bold = True
                r2.font.color.rgb = RGBColor(0xFF, 0x79, 0x00)
                r3 = p1.add_run()
                r3.text = dur_display
                r3.font.name = "Calibri"
                r3.font.size = Pt(9.0)
                r3.font.bold = False
                r3.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

            # Limpiar TextBox 9 si existe para evitar textos residuales
            elif shape.shape_id == 9 and shape.has_text_frame:
                shape.text_frame.clear()

            # 3. IMPACTO (shape_id 11 - CuadroTexto 10)
            elif shape.shape_id == 11:
                shape.left = Inches(0.40)
                shape.top = Inches(0.91)
                shape.width = Inches(9.20)
                shape.height = Inches(0.58)

                tf = shape.text_frame
                tf.word_wrap = True
                tf.margin_left = Pt(2)
                tf.margin_right = Pt(2)
                tf.margin_top = Pt(1)
                tf.margin_bottom = Pt(1)
                tf.clear()

                main_impact, biz_summary = format_impact_content(data.impact_text, data.business_impact)

                total_chars = len(main_impact) + len(biz_summary)
                font_size = Pt(7.5) if total_chars > 360 else (Pt(8.0) if total_chars > 260 else Pt(8.5))

                p1 = tf.paragraphs[0]
                p1.space_after = Pt(2) if biz_summary else Pt(0)
                p1.line_spacing = 1.05
                run1 = p1.add_run()
                run1.text = main_impact
                run1.font.name = "Calibri"
                run1.font.size = font_size
                run1.font.bold = False
                run1.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

                if biz_summary:
                    p2 = tf.add_paragraph()
                    p2.space_after = Pt(0)
                    p2.line_spacing = 1.05
                    label_match = re.match(r'^(Impacto\s+en\s+Negocio\s*:\s*)(.*)', biz_summary, flags=re.I)
                    if label_match:
                        lbl_run = p2.add_run()
                        lbl_run.text = label_match.group(1)
                        lbl_run.font.name = "Calibri"
                        lbl_run.font.size = font_size
                        lbl_run.font.bold = True
                        lbl_run.font.color.rgb = RGBColor(0xFF, 0x79, 0x00)

                        val_run = p2.add_run()
                        val_run.text = label_match.group(2)
                        val_run.font.name = "Calibri"
                        val_run.font.size = font_size
                        val_run.font.bold = False
                        val_run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
                    else:
                        r2 = p2.add_run()
                        r2.text = biz_summary
                        r2.font.name = "Calibri"
                        r2.font.size = font_size
                        r2.font.bold = False
                        r2.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

            # 4. CAUSA (shape_id 22 - CuadroTexto 21)
            elif shape.shape_id == 22:
                shape.left = Inches(0.40)
                shape.top = Inches(1.95)
                shape.width = Inches(4.70)
                shape.height = Inches(1.25)

                tf = shape.text_frame
                tf.word_wrap = True
                tf.margin_left = Pt(2)
                tf.margin_right = Pt(2)
                tf.margin_top = Pt(1)
                tf.margin_bottom = Pt(1)
                tf.clear()

                txt = data.cause_text or "Detalle de causa raíz no disponible."
                font_size = Pt(7.5) if len(txt) > 300 else Pt(8.5)

                p = tf.paragraphs[0]
                p.space_after = Pt(2)
                p.line_spacing = 1.05
                run = p.add_run()
                run.text = txt
                run.font.name = "Calibri"
                run.font.size = font_size
                run.font.bold = False
                run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

            # 5. SOLUCION (shape_id 23 - CuadroTexto 22)
            elif shape.shape_id == 23:
                shape.left = Inches(5.25)
                shape.top = Inches(1.95)
                shape.width = Inches(4.45)
                shape.height = Inches(1.25)

                tf = shape.text_frame
                tf.word_wrap = True
                tf.margin_left = Pt(2)
                tf.margin_right = Pt(2)
                tf.margin_top = Pt(1)
                tf.margin_bottom = Pt(1)
                tf.clear()

                txt = data.solution_text or "Detalle de solución aplicada no disponible."
                font_size = Pt(7.5) if len(txt) > 300 else Pt(8.5)

                p = tf.paragraphs[0]
                p.space_after = Pt(2)
                p.line_spacing = 1.05
                run = p.add_run()
                run.text = txt
                run.font.name = "Calibri"
                run.font.size = font_size
                run.font.bold = False
                run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

            # 6. TABLA DE PUNTOS DE ACCION (4 columnas: Pain Point | Description | Owner | Forecast)
            elif shape.has_table and len(shape.table.columns) == 4:
                table = shape.table
                actions = data.action_points
                target_count = len(actions)

                # Ajustar cantidad de filas de la tabla si hay menos o más
                while len(table.rows) - 1 < target_count:
                    sample_tr = deepcopy(table.rows[len(table.rows) - 1]._tr)
                    table._tbl.append(sample_tr)

                while len(table.rows) - 1 > target_count and len(table.rows) > 2:
                    last_tr = table.rows[len(table.rows) - 1]._tr
                    last_tr.getparent().remove(last_tr)

                if not actions and len(table.rows) > 1:
                    _set_cell_text(table.rows[1].cells[0], "—")
                    _set_cell_text(table.rows[1].cells[1], "Sin puntos de acción registrados")
                    _set_cell_text(table.rows[1].cells[2], "—")
                    _set_cell_text(table.rows[1].cells[3], "—")
                else:
                    for i, ap in enumerate(actions):
                        r_idx = i + 1
                        if r_idx < len(table.rows):
                            _set_cell_text(table.rows[r_idx].cells[0], ap.pain_point or "Otros", font_size_pt=8.0)
                            _set_cell_text(table.rows[r_idx].cells[1], ap.description or "", font_size_pt=8.0)
                            _set_cell_text(table.rows[r_idx].cells[2], ap.owner or "—", font_size_pt=8.0)
                            _set_cell_text(table.rows[r_idx].cells[3], ap.forecast or "—", font_size_pt=8.0)

    def _populate_timeline_slides(self, prs, events: List[ExecutiveTimelineEvent], title_text: str):
        """
        Rellena las diapositivas de cronología distribuyendo los eventos.
        Si no hay eventos, elimina las diapositivas secundarias para dejar un resumen ejecutivo limpio de 1 diapositiva.
        Si hay eventos, genera las diapositivas necesarias en bloques de 18 filas.
        """
        if not events:
            while len(prs.slides) > 1:
                self._delete_slide(prs, 1)
            return

        chunk_size = 18
        chunks: List[List[ExecutiveTimelineEvent]] = []
        for i in range(0, len(events), chunk_size):
            chunks.append(events[i:i + chunk_size])

        needed_total_slides = 1 + len(chunks)
        while len(prs.slides) > needed_total_slides:
            self._delete_slide(prs, len(prs.slides) - 1)

        while len(prs.slides) < needed_total_slides:
            slide_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
            new_slide = prs.slides.add_slide(slide_layout)
            ref_slide = prs.slides[1]
            for shape in ref_slide.shapes:
                if shape.has_table:
                    t = shape.table
                    new_table_shape = new_slide.shapes.add_table(len(t.rows), len(t.columns), shape.left, shape.top, shape.width, shape.height)
                    for r_idx in range(len(t.rows)):
                        for c_idx in range(len(t.columns)):
                            new_table_shape.table.cell(r_idx, c_idx).text = t.cell(r_idx, c_idx).text
                elif shape.has_text_frame:
                    new_tx = new_slide.shapes.add_textbox(shape.left, shape.top, shape.width, shape.height)
                    new_tx.text_frame.text = shape.text_frame.text

        for idx, chunk in enumerate(chunks):
            slide_idx = idx + 1
            if slide_idx >= len(prs.slides):
                break
            slide = prs.slides[slide_idx]
            for shape in slide.shapes:
                if shape.shape_id == 4 or (shape.has_text_frame and "RESUMEN EJECUTIVO INCIDENCIA" in shape.text_frame.text):
                    tf = shape.text_frame
                    tf.clear()
                    p = tf.paragraphs[0]
                    run = p.add_run()
                    run.text = title_text
                    run.font.name = "Calibri"
                    run.font.bold = True
                    run.font.size = Pt(10.0) if len(title_text) > 80 else Pt(11.5)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                elif shape.has_table and len(shape.table.columns) == 2:
                    table = shape.table
                    target_rows = max(len(chunk), 1)
                    while len(table.rows) - 1 < target_rows:
                        sample_tr = deepcopy(table.rows[len(table.rows) - 1]._tr)
                        table._tbl.append(sample_tr)
                    while len(table.rows) - 1 > target_rows and len(table.rows) > 2:
                        last_tr = table.rows[len(table.rows) - 1]._tr
                        last_tr.getparent().remove(last_tr)

                    for e_idx, event in enumerate(chunk):
                        r_idx = e_idx + 1
                        if r_idx < len(table.rows):
                            _set_cell_text(table.rows[r_idx].cells[0], event.time or "", font_size_pt=8.0)
                            _set_cell_text(table.rows[r_idx].cells[1], event.event or "", font_size_pt=8.0)
